from openpyxl import load_workbook

wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
# bar_value = Reference(ws,min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart() # 차트 종류 설정(Bar, Line, Pie...)
# bar_chart.add_data(bar_value)

# ws.add_chart(bar_chart, "E1") # 차트 넣을 위치 정의

# B1:C11 가지의 데이터
line_value = Reference(ws,min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart() 
line_chart.add_data(line_value, titles_from_data=True) # 계열이 제목에서 가져오기
line_chart.title="성적표"
line_chart.style  =20 # 미리 정의된 스타일을 적용
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"

ws.add_chart(line_chart, "E1") # 차트 넣을 위치 정의

wb.save("sample_chart.xlsx")